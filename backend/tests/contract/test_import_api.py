"""T051：产品与导入接口契约测试（contracts/admin-api.md）。

上传接口是异步的（research.md R3）：POST 只登记批次并返回 202，
真正的解析由后台 worker 完成。测试里直接同步调用 process_batch，
避免依赖线程调度。
"""

import io

import pytest
from openpyxl import Workbook, load_workbook
from sqlalchemy.orm import Session

from app.models.import_batch import ImportBatch
from app.models.user import User
from app.tasks import import_worker
from tests.conftest import auth_header

pytestmark = pytest.mark.contract


def _xlsx_bytes(header: list[str], rows: list[list[str]]) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(header)
    for row in rows:
        sheet.append(row)
    bio = io.BytesIO()
    workbook.save(bio)
    return bio.getvalue()


def _upload(client, user, codes: list[str], mapping: list[tuple[str, str]]):
    return client.post(
        "/api/v1/imports",
        headers=auth_header(user),
        files={
            "code_list_file": (
                "codes.xlsx",
                _xlsx_bytes(["码值"], [[c] for c in codes]),
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            ),
            "mapping_file": (
                "mapping.xlsx",
                _xlsx_bytes(["码值", "产品编码"], [[c, p] for c, p in mapping]),
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            ),
        },
    )


# ---------------------------------------------------------------- 产品


class TestProductEndpoints:
    def test_create_and_list_with_stats(self, client, hq_admin: User):
        created = client.post(
            "/api/v1/products",
            headers=auth_header(hq_admin),
            json={"code": "PX01", "name": "测试产品", "spec": "1L"},
        )
        assert created.status_code == 201
        assert created.json()["status"] == "active"
        assert created.json()["stats"] == {"imported": 0, "allocated": 0, "printed": 0}

        listed = client.get("/api/v1/products", headers=auth_header(hq_admin)).json()
        assert any(p["code"] == "PX01" for p in listed)

    def test_duplicate_product_code_rejected(self, client, hq_admin, product):
        resp = client.post(
            "/api/v1/products",
            headers=auth_header(hq_admin),
            json={"code": product.code, "name": "重名"},
        )
        assert resp.status_code == 422
        assert resp.json()["error_code"] == "VALIDATION_ERROR"

    def test_search_by_keyword(self, client, hq_admin, product):
        body = client.get(
            f"/api/v1/products?keyword={product.name[:3]}", headers=auth_header(hq_admin)
        ).json()
        assert all(product.name[:3] in p["name"] or product.code in p["code"] for p in body)

    def test_sub_admin_cannot_create_product(self, client, sub_admin):
        resp = client.post(
            "/api/v1/products",
            headers=auth_header(sub_admin),
            json={"code": "PZ01", "name": "越权产品"},
        )
        assert resp.status_code == 403
        assert resp.json()["error_code"] == "FORBIDDEN_SCOPE"

    def test_sub_admin_can_read_products(self, client, sub_admin, product):
        resp = client.get("/api/v1/products", headers=auth_header(sub_admin))
        assert resp.status_code == 200


# ---------------------------------------------------------------- 导入


class TestImportTemplate:
    def test_template_has_both_sheets(self, client, hq_admin: User):
        resp = client.get("/api/v1/imports/template", headers=auth_header(hq_admin))
        assert resp.status_code == 200
        assert "spreadsheetml" in resp.headers["content-type"]

        workbook = load_workbook(io.BytesIO(resp.content))
        assert workbook.sheetnames == ["码清单", "码-产品对照关系"]
        assert workbook["码清单"]["A1"].value == "码值"
        assert [c.value for c in workbook["码-产品对照关系"][1]] == ["码值", "产品编码"]

    def test_template_requires_hq(self, client, sub_admin):
        resp = client.get("/api/v1/imports/template", headers=auth_header(sub_admin))
        assert resp.status_code == 403


class TestImportUpload:
    def test_upload_returns_202_with_batch(
        self, client, hq_admin: User, product, import_storage, db: Session
    ):
        resp = _upload(client, hq_admin, ["U001"], [("U001", product.code)])
        assert resp.status_code == 202

        body = resp.json()
        assert set(body) == {"batch_id", "batch_no", "status"}
        assert body["status"] == "pending"
        assert body["batch_no"].startswith("IMP")

        batch = db.get(ImportBatch, body["batch_id"])
        assert batch is not None
        assert batch.code_list_filename == "codes.xlsx"
        assert batch.mapping_filename == "mapping.xlsx"

    def test_upload_does_not_parse_synchronously(
        self, client, hq_admin: User, product, import_storage, db: Session
    ):
        """R3：上传阶段不得解析文件，成功数必须还是 0。"""
        from app.models.code import Code

        resp = _upload(client, hq_admin, ["U002"], [("U002", product.code)])
        batch = db.get(ImportBatch, resp.json()["batch_id"])

        assert batch.success_count == 0
        assert batch.processed_rows == 0
        assert db.query(Code).count() == 0

    def test_worker_then_completes_the_batch(
        self, client, hq_admin: User, product, import_storage, db: Session
    ):
        resp = _upload(
            client, hq_admin, ["U003", "U004"], [("U003", product.code), ("U004", product.code)]
        )
        batch_id = resp.json()["batch_id"]

        import_worker.process_batch(db, batch_id)

        detail = client.get(f"/api/v1/imports/{batch_id}", headers=auth_header(hq_admin)).json()
        assert detail["status"] == "succeeded"
        assert detail["success_count"] == 2
        assert detail["processed_rows"] == 2

    def test_missing_file_is_rejected(self, client, hq_admin, import_storage):
        resp = client.post(
            "/api/v1/imports",
            headers=auth_header(hq_admin),
            files={
                "code_list_file": ("codes.xlsx", _xlsx_bytes(["码值"], [["X1"]]), "application/x"),
            },
        )
        assert resp.status_code == 422

    def test_sub_admin_cannot_import(self, client, sub_admin, import_storage):
        resp = _upload(client, sub_admin, ["U005"], [("U005", "P001")])
        assert resp.status_code == 403


class TestImportQuery:
    def test_batch_detail_shape(
        self, client, hq_admin: User, product, import_storage, db: Session
    ):
        resp = _upload(client, hq_admin, ["Q001"], [("Q001", product.code)])
        batch_id = resp.json()["batch_id"]

        body = client.get(f"/api/v1/imports/{batch_id}", headers=auth_header(hq_admin)).json()
        assert set(body) >= {
            "id",
            "batch_no",
            "status",
            "total_rows",
            "processed_rows",
            "success_count",
            "failure_count",
            "created_at",
        }

    def test_unknown_batch_is_404(self, client, hq_admin: User):
        resp = client.get("/api/v1/imports/999999", headers=auth_header(hq_admin))
        assert resp.status_code == 404

    def test_batch_list_returns_newest_first(
        self, client, hq_admin: User, product, import_storage
    ):
        first = _upload(client, hq_admin, ["L001"], [("L001", product.code)]).json()["batch_id"]
        second = _upload(client, hq_admin, ["L002"], [("L002", product.code)]).json()["batch_id"]

        body = client.get("/api/v1/imports", headers=auth_header(hq_admin)).json()
        ids = [b["id"] for b in body]
        assert ids.index(second) < ids.index(first)

    def test_errors_download_404_when_none(
        self, client, hq_admin: User, product, import_storage
    ):
        resp = _upload(client, hq_admin, ["E001"], [("E001", product.code)])
        batch_id = resp.json()["batch_id"]
        assert (
            client.get(f"/api/v1/imports/{batch_id}/errors", headers=auth_header(hq_admin)).status_code
            == 404
        )

    def test_errors_download_returns_csv(
        self, client, hq_admin: User, product, import_storage, db: Session
    ):
        # 对照关系引用不存在的产品 → 产生失败明细
        resp = _upload(client, hq_admin, ["E002"], [("E002", "NO-SUCH")])
        batch_id = resp.json()["batch_id"]
        import_worker.process_batch(db, batch_id)

        download = client.get(
            f"/api/v1/imports/{batch_id}/errors", headers=auth_header(hq_admin)
        )
        assert download.status_code == 200
        text = download.content.decode("utf-8-sig")
        assert "行号" in text
        assert "PRODUCT_NOT_FOUND" in text
