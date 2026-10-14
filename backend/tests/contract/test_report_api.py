"""T131：报表接口契约测试（contracts/admin-api.md 报表与审计小节）。"""

from datetime import timedelta

import pytest
from sqlalchemy.orm import Session

from app.models.base import utcnow
from app.models.user import User
from tests.conftest import auth_header
from tests.helpers import make_printed

pytestmark = pytest.mark.contract


@pytest.fixture
def printed(db: Session, product, hq_company, import_batch, template_version, hq_admin, workstation, operator):
    make_printed(
        db,
        product=product,
        batch=import_batch,
        company=hq_company,
        operator_id=hq_admin.id,
        workstation=workstation,
        values=["RP01", "RP02"],
    )
    return ["RP01", "RP02"]


class TestDashboardEndpoint:
    def test_response_shape(self, client, hq_admin: User, printed):
        resp = client.get(
            "/api/v1/dashboard/code-distribution", headers=auth_header(hq_admin)
        )
        assert resp.status_code == 200

        body = resp.json()
        assert set(body) == {"group_by", "rows", "totals"}
        assert set(body["totals"]) == {
            "key",
            "company_name",
            "product_name",
            "unallocated",
            "allocated",
            "printing",
            "printed",
            "voided",
            "allocated_not_printed",
            "total",
        }
        assert body["totals"]["printed"] == 2

    def test_requires_authentication(self, client):
        assert client.get("/api/v1/dashboard/code-distribution").status_code == 401

    def test_sub_admin_can_access(self, client, sub_admin: User, printed):
        assert (
            client.get(
                "/api/v1/dashboard/code-distribution", headers=auth_header(sub_admin)
            ).status_code
            == 200
        )


class TestPrintDetailsEndpoint:
    def test_response_shape(self, client, hq_admin: User, printed):
        rows = client.get("/api/v1/reports/prints", headers=auth_header(hq_admin)).json()
        assert rows
        assert set(rows[0]) == {
            "code_value",
            "product_name",
            "company_name",
            "workstation_name",
            "operator_name",
            "printer_name",
            "print_type",
            "result",
            "failure_reason",
            "printed_at",
        }
        assert rows[0]["print_type"] == "first"
        assert rows[0]["result"] == "succeeded"

    def test_filters_by_workstation(self, client, hq_admin, workstation, printed):
        rows = client.get(
            f"/api/v1/reports/prints?workstation_id={workstation.id}",
            headers=auth_header(hq_admin),
        ).json()
        assert len(rows) == 2

        empty = client.get(
            "/api/v1/reports/prints?workstation_id=999999", headers=auth_header(hq_admin)
        ).json()
        assert empty == []

    def test_filters_by_operator(self, client, hq_admin, printed):
        # 夹具里的打印动作是以 hq_admin 身份执行的，记录也就归在它名下
        rows = client.get(
            f"/api/v1/reports/prints?operator_id={hq_admin.id}",
            headers=auth_header(hq_admin),
        ).json()
        assert len(rows) == 2

        other = client.get(
            "/api/v1/reports/prints?operator_id=999999", headers=auth_header(hq_admin)
        ).json()
        assert other == []

    def test_filters_by_product(self, client, hq_admin, product, printed):
        rows = client.get(
            f"/api/v1/reports/prints?product_id={product.id}", headers=auth_header(hq_admin)
        ).json()
        assert len(rows) == 2

    def test_filters_by_time_window(self, client, hq_admin, printed):
        future = (utcnow() + timedelta(days=1)).strftime("%Y-%m-%dT%H:%M:%S")
        rows = client.get(
            f"/api/v1/reports/prints?printed_from={future}", headers=auth_header(hq_admin)
        ).json()
        assert rows == []

        past = (utcnow() - timedelta(days=1)).strftime("%Y-%m-%dT%H:%M:%S")
        rows = client.get(
            f"/api/v1/reports/prints?printed_to={past}", headers=auth_header(hq_admin)
        ).json()
        assert rows == []

    def test_export_returns_xlsx(self, client, hq_admin: User, printed):
        resp = client.get("/api/v1/reports/prints/export", headers=auth_header(hq_admin))
        assert resp.status_code == 200
        assert "spreadsheetml" in resp.headers["content-type"]
        assert "attachment" in resp.headers["content-disposition"]
        # xlsx 是 zip 容器，前两字节固定为 PK
        assert resp.content[:2] == b"PK"

    def test_export_headers_are_readable(self, client, hq_admin: User, printed):
        import io

        from openpyxl import load_workbook

        resp = client.get("/api/v1/reports/prints/export", headers=auth_header(hq_admin))
        workbook = load_workbook(io.BytesIO(resp.content))
        sheet = workbook.active
        headers = [c.value for c in sheet[1]]
        assert headers[0] == "码值"
        assert "打印时间" in headers
        assert sheet.max_row == 3  # 表头 + 2 条记录

    def test_export_respects_filters(self, client, hq_admin: User, printed):
        import io

        from openpyxl import load_workbook

        resp = client.get(
            "/api/v1/reports/prints/export?workstation_id=999999",
            headers=auth_header(hq_admin),
        )
        workbook = load_workbook(io.BytesIO(resp.content))
        assert workbook.active.max_row == 1  # 只剩表头

    def test_requires_authentication(self, client):
        assert client.get("/api/v1/reports/prints").status_code == 401
