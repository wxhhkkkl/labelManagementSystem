"""T056：流式 Excel 解析。

MUST 使用 openpyxl 的 read_only 模式逐行读取（research.md R3）：
一次性 load_workbook 会把 10 万行全部载入内存，占用数百 MB。
"""

from collections.abc import Iterator
from pathlib import Path

from openpyxl import load_workbook

CODE_LIST_HEADER = "码值"
MAPPING_HEADER = ("码值", "产品编码")


class ImportFileError(Exception):
    """文件本身不可解析（空文件、缺列等），与"行内容不合法"区分开。"""


def _first_column(row: tuple) -> object:
    return row[0] if row else None


def _cell(row: tuple, index: int) -> object:
    return row[index] if row and len(row) > index else None


def _open(path: Path):
    if not path.exists():
        raise ImportFileError(f"文件不存在：{path.name}")
    try:
        workbook = load_workbook(path, read_only=True, data_only=True)
    except Exception as exc:  # noqa: BLE001
        raise ImportFileError(f"无法解析 {path.name}：{exc}") from exc
    sheet = workbook.active
    if sheet is None:
        workbook.close()
        raise ImportFileError(f"{path.name} 中没有工作表")
    return workbook, sheet


def iter_code_list(path: Path) -> Iterator[tuple[int, object]]:
    """逐行产出 (Excel 行号, 首列原始值)。第 1 行视为表头跳过。"""
    workbook, sheet = _open(path)
    try:
        for row_no, row in enumerate(sheet.iter_rows(min_row=2, values_only=True), start=2):
            yield row_no, _first_column(row)
    finally:
        workbook.close()


def iter_mapping(path: Path) -> Iterator[tuple[int, object, object]]:
    """逐行产出 (Excel 行号, 码值, 产品编码)。第 1 行视为表头跳过。"""
    workbook, sheet = _open(path)
    try:
        for row_no, row in enumerate(sheet.iter_rows(min_row=2, values_only=True), start=2):
            yield row_no, _cell(row, 0), _cell(row, 1)
    finally:
        workbook.close()


def iter_product_batch(path: Path) -> Iterator[tuple[int, object, object, str | None]]:
    """Read named columns; external query status is deliberately not a print state."""
    try:
        workbook = load_workbook(path, read_only=True, data_only=False)
    except Exception as exc:
        raise ImportFileError(f"无法解析 {path.name}：{exc}") from exc
    try:
        sheet = workbook.active
        if sheet is None:
            raise ImportFileError("文件没有工作表")
        header = next(sheet.iter_rows(min_row=1, max_row=1, values_only=True), ())
        names = [str(value or "").strip() for value in header]
        for name in ("防伪码", "产品批号"):
            if names.count(name) != 1:
                raise ImportFileError(f"首行必须包含唯一的“{name}”列")
        code_index, batch_index = names.index("防伪码"), names.index("产品批号")
        for row_no, row in enumerate(sheet.iter_rows(min_row=2), 2):
            if all(cell.value is None for cell in row):
                continue
            code_cell, batch_cell = row[code_index], row[batch_index]
            error = None
            if code_cell.data_type == "f" or batch_cell.data_type == "f":
                error = "防伪码和产品批号不能使用公式，请粘贴为文本值"
            elif (
                isinstance(code_cell.value, (int, float))
                and len(str(abs(int(code_cell.value)))) > 15
            ):
                error = "超过15位的防伪码必须为文本，数字单元格可能已丢失精度"
            yield row_no, code_cell.value, batch_cell.value, error
    finally:
        workbook.close()
