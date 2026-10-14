"""导入相关测试的公共构造工具。"""

from pathlib import Path

from openpyxl import Workbook
from sqlalchemy.orm import Session

from app.models.import_batch import ImportBatch
from app.tasks import import_worker


def write_code_list(path: Path, values: list[str]) -> Path:
    """写码清单文件（首列为码值）。"""
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook = Workbook(write_only=True)
    sheet = workbook.create_sheet("码清单")
    sheet.append(["码值"])
    for value in values:
        sheet.append([value])
    workbook.save(path)
    return path


def write_mapping(path: Path, rows: list[tuple[str, str]]) -> Path:
    """写"码—产品"对照关系文件（码值、产品编码）。"""
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook = Workbook(write_only=True)
    sheet = workbook.create_sheet("对照关系")
    sheet.append(["码值", "产品编码"])
    for code, product in rows:
        sheet.append([code, product])
    workbook.save(path)
    return path


def stage_batch(
    db: Session,
    *,
    batch_no: str,
    created_by: int,
    code_values: list[str],
    mapping_rows: list[tuple[str, str]],
) -> ImportBatch:
    """把两个文件写到 worker 约定的位置，并登记批次。"""
    write_code_list(import_worker.code_list_path(batch_no), code_values)
    write_mapping(import_worker.mapping_path(batch_no), mapping_rows)

    batch = ImportBatch(
        batch_no=batch_no,
        code_list_filename="code_list.xlsx",
        mapping_filename="mapping.xlsx",
        status="pending",
        processed_rows=0,
        success_count=0,
        failure_count=0,
        created_by=created_by,
    )
    db.add(batch)
    db.flush()
    return batch


def paired_rows(codes: list[str], product_code: str) -> tuple[list[str], list[tuple[str, str]]]:
    """生成两个文件的内容：码清单与一一对应的对照关系。"""
    return codes, [(code, product_code) for code in codes]


def make_printable(
    db: Session,
    *,
    product,
    batch,
    company,
    operator_id: int,
    values: list[str],
):
    """把一批码走完整链路推到"已分配"，供打印相关测试使用。

    走真实的 allocation_service 而不是直接写 status='allocated'：
    打印时的模板解析依赖分配单携带的模板版本，绕过去就测不到真实路径。
    """
    from app.models.base import utcnow
    from app.models.code import Code
    from app.models.code_event import CodeEvent
    from app.services import allocation as allocation_service

    codes = [
        Code(
            code_value=value,
            status="unallocated",
            product_id=product.id,
            import_batch_id=batch.id,
        )
        for value in values
    ]
    db.add_all(codes)
    db.flush()

    # 真实的码都来自导入流程，必然带一条 imported 事件。
    # 这里直接插码会造出"库里有码但没导入记录"的假状态，
    # 时间线之类依赖事件链的测试会因此失真。
    now = utcnow()
    db.add_all(
        [
            CodeEvent(
                code_id=code.id,
                event_type="imported",
                from_status=None,
                to_status="unallocated",
                operator_id=operator_id,
                occurred_at=now,
            )
            for code in codes
        ]
    )
    db.flush()

    result = allocation_service.allocate(
        db,
        company_id=company.id,
        product_id=product.id,
        quantity=len(values),
        operator_id=operator_id,
    )
    db.flush()
    return result.order


def make_printed(
    db: Session,
    *,
    product,
    batch,
    company,
    operator_id: int,
    workstation,
    values: list[str],
):
    """把一批码推到"已打印"，供补打相关测试使用。

    走真实的预占 + 确认链路，这样 last_printed_workstation_id 与打印记录
    都是真实产生的——补打派发正是依赖前者。
    """
    from app.services import printing

    make_printable(
        db,
        product=product,
        batch=batch,
        company=company,
        operator_id=operator_id,
        values=values,
    )
    printing.preempt(
        db, workstation=workstation, operator_id=operator_id, code_values=values
    )
    printing.confirm(
        db,
        workstation=workstation,
        operator_id=operator_id,
        results=[
            {"code_value": value, "outcome": "succeeded", "printer_name": "Zebra ZD421"}
            for value in values
        ],
    )
    db.flush()
